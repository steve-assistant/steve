"""Google Drive, Docs, Sheets and Slides.

The creation strategy is the interesting decision here, and it is the opposite
of what the API documentation suggests.

`documents.batchUpdate` operates on a flat index space where every mutation
invalidates the indices of everything after it. Producing a formatted document
means dozens of requests ordered by *descending* index, with a refetch between
structural operations -- and the failure mode when a model gets it wrong is
not an error but silent corruption: text lands in the wrong paragraph and
nothing complains.

So documents are built locally with python-docx / python-pptx and uploaded
with `mimeType: application/vnd.google-apps.*`, which converts them to native
Google files on the way in. One call, full formatting, no index arithmetic.

Sheets is the exception. `values.update` takes A1 notation and a 2D array;
there is no index space to get wrong, so writing cells natively is both
simpler and keeps the file live.
"""

from __future__ import annotations

import io
from typing import Any

from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

from .common import clamp, guard, service

GOOGLE_DOC = "application/vnd.google-apps.document"
GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"
GOOGLE_SLIDES = "application/vnd.google-apps.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Exporting a Google-native file needs a target format; these are the sane
# defaults for reading one back as text.
EXPORT_AS = {
    GOOGLE_DOC: "text/plain",
    GOOGLE_SHEET: "text/csv",
    GOOGLE_SLIDES: "text/plain",
}

MAX_TEXT_CHARS = 20_000


def _drive():
    return service("drive", "v3")


def _sheets():
    return service("sheets", "v4")


def _file_info(entry: dict) -> dict:
    return {
        "id": entry.get("id"),
        "name": entry.get("name"),
        "mime_type": entry.get("mimeType"),
        "modified": entry.get("modifiedTime"),
        "size": entry.get("size"),
        "link": entry.get("webViewLink"),
        "owner": (entry.get("owners") or [{}])[0].get("emailAddress"),
    }


FIELDS = "files(id,name,mimeType,modifiedTime,size,webViewLink,owners(emailAddress))"


def _files(query: str, limit: int) -> list[dict]:
    """The files a Drive query finds, most recently modified first."""
    found = _drive().files().list(q=query, pageSize=limit, fields=FIELDS,
                                  orderBy="modifiedTime desc").execute()
    return [_file_info(f) for f in found.get("files", [])]


@guard
def search(query: str = "", *, limit: int = 20, mime_type: str | None = None) -> dict:
    """Find files by name or content.

    `query` is matched against both the name and full text. Drive's own query
    language is not exposed, because its quoting rules are a trap -- an
    unescaped apostrophe in a filename produces a syntax error, not no results.
    """
    clauses = ["trashed = false"]
    if query:
        safe = query.replace("\\", "\\\\").replace("'", "\\'")
        clauses.append(f"(name contains '{safe}' or fullText contains '{safe}')")
    if mime_type:
        clauses.append(f"mimeType = '{mime_type}'")

    files = _files(" and ".join(clauses), clamp(limit, 20))
    return {"query": query, "count": len(files), "files": files}


@guard
def recent(limit: int = 15) -> dict:
    """Most recently modified files."""
    return {"files": _files("trashed = false", clamp(limit, 15))}


@guard
def read(file_id: str, *, max_chars: int = MAX_TEXT_CHARS) -> dict:
    """Read a file's text.

    Google-native files have no bytes to download and must be exported;
    everything else is fetched directly. Binary formats return a description
    rather than mojibake.
    """
    meta = _drive().files().get(fileId=file_id, fields="id,name,mimeType,size").execute()
    mime = meta.get("mimeType", "")

    if mime in EXPORT_AS:
        data = _drive().files().export(fileId=file_id, mimeType=EXPORT_AS[mime]).execute()
        text = data.decode("utf-8", "replace") if isinstance(data, bytes) else str(data)
    elif mime.startswith("text/") or mime in ("application/json", "application/xml"):
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, _drive().files().get_media(fileId=file_id))
        done = False
        while not done:
            _, done = downloader.next_chunk()
        text = buffer.getvalue().decode("utf-8", "replace")
    else:
        return {
            "id": file_id,
            "name": meta.get("name"),
            "mime_type": mime,
            "error": f"cannot read {mime} as text; download it instead",
            "kind": "unsupported",
        }

    return {
        "id": file_id,
        "name": meta.get("name"),
        "mime_type": mime,
        "text": text[:max_chars],
        "truncated": len(text) > max_chars,
    }


def _upload_converted(
    data: bytes, source_mime: str, name: str, target_mime: str, folder_id: str | None
) -> dict:
    metadata: dict[str, Any] = {"name": name, "mimeType": target_mime}
    if folder_id:
        metadata["parents"] = [folder_id]
    media = MediaIoBaseUpload(io.BytesIO(data), mimetype=source_mime, resumable=False)
    created = (
        _drive().files()
        .create(body=metadata, media_body=media, fields="id,name,mimeType,webViewLink")
        .execute()
    )
    return _file_info(created)


@guard
def create_doc(
    title: str,
    content: str,
    *,
    folder_id: str | None = None,
    heading: str | None = None,
) -> dict:
    """Create a Google Doc from markdown-ish text.

    Built as a .docx and converted on upload rather than assembled with
    documents.batchUpdate -- see the module docstring for why that API is a
    trap for a model.
    """
    from docx import Document

    document = Document()
    if heading:
        document.add_heading(heading, level=1)

    for block in content.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("### "):
            document.add_heading(block[4:], level=3)
        elif block.startswith("## "):
            document.add_heading(block[3:], level=2)
        elif block.startswith("# "):
            document.add_heading(block[2:], level=1)
        elif all(line.lstrip().startswith(("- ", "* ")) for line in block.splitlines()):
            for line in block.splitlines():
                document.add_paragraph(line.lstrip()[2:], style="List Bullet")
        else:
            document.add_paragraph(block)

    # Built in memory. It went through a folder in %TEMP% until step 5's
    # item 4f: every program of this Windows user shares that folder, and
    # removing Steve's data would miss what a failed upload left there.
    built = io.BytesIO()
    document.save(built)
    info = _upload_converted(built.getvalue(), DOCX, title, GOOGLE_DOC, folder_id)
    return {"created": True, **info}


@guard
def create_sheet(
    title: str,
    rows: list[list[Any]],
    *,
    sheet_name: str = "Sheet1",
    folder_id: str | None = None,
) -> dict:
    """Create a Google Sheet and populate it.

    Native here, unlike Docs: `values.update` takes A1 notation and a 2D
    array, so there is no index arithmetic to get wrong, and the result stays
    a live spreadsheet rather than a converted upload.
    """
    created = (
        _sheets().spreadsheets()
        .create(
            body={
                "properties": {"title": title},
                "sheets": [{"properties": {"title": sheet_name}}],
            },
            fields="spreadsheetId,spreadsheetUrl",
        )
        .execute()
    )
    spreadsheet_id = created["spreadsheetId"]

    if rows:
        _sheets().spreadsheets().values().update(
            spreadsheetId=spreadsheet_id,
            range=f"{sheet_name}!A1",
            # USER_ENTERED so "=SUM(A1:A9)" becomes a formula and "50%"
            # becomes a percentage, exactly as if typed.
            valueInputOption="USER_ENTERED",
            body={"values": rows},
        ).execute()

    if folder_id:
        _drive().files().update(
            fileId=spreadsheet_id, addParents=folder_id, fields="id"
        ).execute()

    return {
        "created": True,
        "id": spreadsheet_id,
        "name": title,
        "link": created.get("spreadsheetUrl"),
        "rows_written": len(rows),
    }


@guard
def read_sheet(spreadsheet_id: str, *, range_name: str = "A1:Z1000") -> dict:
    """Read cell values from a spreadsheet."""
    result = (
        _sheets().spreadsheets().values()
        .get(spreadsheetId=spreadsheet_id, range=range_name)
        .execute()
    )
    values = result.get("values", [])
    return {"id": spreadsheet_id, "range": result.get("range"), "rows": len(values),
            "values": values}


@guard
def update_sheet(
    spreadsheet_id: str, rows: list[list[Any]], *, range_name: str = "A1"
) -> dict:
    """Write cells into an existing spreadsheet."""
    result = (
        _sheets().spreadsheets().values()
        .update(
            spreadsheetId=spreadsheet_id,
            range=range_name,
            valueInputOption="USER_ENTERED",
            body={"values": rows},
        )
        .execute()
    )
    return {
        "updated": True,
        "id": spreadsheet_id,
        "cells": result.get("updatedCells"),
        "range": result.get("updatedRange"),
    }
