"""The world the connectors' golden outputs are recorded in (step 5, item 4a).

Step 5 moves every connector. The mail tools, the calendar, Drive, Tasks and
the attention sources stop calling `tools/*` and go through a slot, and the
code they called moves into `modules/`. A move like that can change what the
user sees while every old test still passes, because each old test fakes the
very function the move replaced. So the fakes here sit below all of it, at
each service's own seam: Google's API client, Outlook's MAPI namespace, the
IMAP and SMTP connections, a calendar link's feed. What is recorded is what the
tools returned, what each would have asked, and what reached the services.
As the code moves, only `wire` may change; the record may not.

Everything else is pinned too: the clock, the zone, and the ids the
services hand out, so that one run is the same as the next.
"""

from __future__ import annotations

import base64
import copy
import dataclasses
import datetime as _dt
import email
import email.message
import email.policy
import email.utils
import inspect
import io
import itertools
import re
import shlex
import sys
import time as _time
import types
from dataclasses import dataclass, field
from zoneinfo import ZoneInfo

import pytest

from tests import persona
from tests.fakes import ics as feeds
from tests.fakes import mailrig
from tests.fakes.imap import STANDARD

UTC = _dt.UTC
ZONE = ZoneInfo(persona.ZONE)
# A Wednesday morning in the user's zone: a standup at 10, a review at 11.
MOMENT = _dt.datetime(2026, 10, 14, 9, 30, tzinfo=ZONE)
# Their Google account's address: a mail account's, so their own files and
# mail read as theirs (`agent.connectors.addresses`).
ME = persona.EMAIL


def at(days: float = 0, hours: float = 0, minutes: float = 0) -> _dt.datetime:
    """A moment relative to the pinned one, in the user's zone."""
    return MOMENT + _dt.timedelta(days=days, hours=hours, minutes=minutes)


# ---------------------------------------------------------------------------
# The clock
# ---------------------------------------------------------------------------


class _Pinned(type):
    """A stand-in for `datetime.datetime` or `datetime.date` whose `now` is
    pinned. Everything else -- construction, `isinstance`, `fromisoformat` --
    is the real class's, so a value made through it is a real datetime."""

    def __call__(cls, *args, **kwargs):
        return cls.real(*args, **kwargs)

    def __instancecheck__(cls, obj) -> bool:
        return isinstance(obj, cls.real)

    def __subclasscheck__(cls, sub) -> bool:
        return issubclass(sub, cls.real)

    def __getattr__(cls, name):
        return getattr(cls.real, name)


def _pinned_classes(moment: _dt.datetime):
    class PinnedDateTime(metaclass=_Pinned):
        real = _dt.datetime

        @staticmethod
        def now(tz=None):
            if tz is None:  # naive: the wall clock of the user's zone
                return moment.astimezone(ZONE).replace(tzinfo=None)
            return moment.astimezone(tz)

        @staticmethod
        def utcnow():
            return moment.astimezone(UTC).replace(tzinfo=None)

        @staticmethod
        def today():
            return moment.astimezone(ZONE).replace(tzinfo=None)

    class PinnedDate(metaclass=_Pinned):
        real = _dt.date

        @staticmethod
        def today():
            return moment.astimezone(ZONE).date()

    return PinnedDateTime, PinnedDate


class _PinnedTime(types.ModuleType):
    """`time`, whose `time()` starts at the pinned moment and runs on from
    there. Running, not stopped: a loop waiting on a deadline would wait
    forever on a clock that never moves."""

    def __init__(self, moment: _dt.datetime) -> None:
        super().__init__("time")
        self._at, self._from = moment.timestamp(), _time.monotonic()

    def time(self) -> float:
        return self._at + (_time.monotonic() - self._from)

    def time_ns(self) -> int:
        return int(self.time() * 1e9)

    def __getattr__(self, name):
        return getattr(_time, name)


class _PinnedModule(types.ModuleType):
    def __init__(self, real: types.ModuleType, **pinned) -> None:
        super().__init__(real.__name__)
        self.__dict__.update(pinned)
        self._real = real

    def __getattr__(self, name):
        return getattr(self._real, name)


# The packages whose clocks are pinned: Steve's own code, wherever it moves.
PACKAGES = ("agent", "tools", "core", "steve_sdk", "modules")


def ours(name: str) -> bool:
    return name.split(".", 1)[0] in PACKAGES


_CLOCKS = {id(_dt.datetime), id(_dt.date), id(_dt), id(_time), id(_time.time)}


def unpinned(names) -> list[str]:
    """Those of these modules that hold a clock of their own: imported after
    `pin_the_clock`, they read the real one."""
    return sorted(n for n in names if (module := sys.modules.get(n)) is not None
                  and any(id(v) in _CLOCKS for v in vars(module).values()))


def pin_the_clock(monkeypatch, moment: _dt.datetime = MOMENT) -> None:
    """Pin `datetime.now`, `date.today` and `time.time` in every module of
    Steve's already imported, however each imported them."""
    pinned_datetime, pinned_date = _pinned_classes(moment)
    pinned_module = _PinnedModule(_dt, datetime=pinned_datetime, date=pinned_date)
    pinned_time = _PinnedTime(moment)
    swaps = {id(_dt.datetime): pinned_datetime, id(_dt.date): pinned_date,
             id(_dt): pinned_module, id(_time): pinned_time, id(_time.time): pinned_time.time}
    for name, module in list(sys.modules.items()):
        if module is None or not ours(name):
            continue
        for attr, value in list(vars(module).items()):
            if (swap := swaps.get(id(value))) is not None:
                monkeypatch.setattr(module, attr, swap)


# ---------------------------------------------------------------------------
# Google: one client for Gmail, Calendar, Drive and Tasks
# ---------------------------------------------------------------------------


def _b64(text: str | bytes) -> str:
    """Base64url as Gmail's API gives a body: padded, which its reader expects."""
    data = text.encode() if isinstance(text, str) else text
    return base64.urlsafe_b64encode(data).decode()


def _unb64(text: str) -> bytes:
    return base64.urlsafe_b64decode(text + "=" * (-len(text) % 4))


def _not_found(what: str):
    import httplib2
    from googleapiclient.errors import HttpError

    return HttpError(httplib2.Response({"status": "404"}),
                     b'{"error": {"code": 404, "message": "Requested entity was not found."}}',
                     uri=what)


def _precondition_failed(what: str):
    import httplib2
    from googleapiclient.errors import HttpError

    return HttpError(httplib2.Response({"status": "412"}),
                     b'{"error": {"code": 412, "message": "Precondition Failed"}}', uri=what)


def _written(raw: str) -> dict:
    """What a raw RFC 822 message says, for the record: who, what, and the
    words, without the date and Message-ID the mail library makes up."""
    message = email.message_from_bytes(_unb64(raw), policy=email.policy.default)
    body = message.get_body(preferencelist=("plain", "html"))
    said = {name: str(message.get(header) or "") for name, header in (
        ("from", "From"), ("to", "To"), ("cc", "Cc"), ("subject", "Subject"),
        ("in_reply_to", "In-Reply-To"), ("references", "References"))}
    said["body"] = body.get_content().strip() if body is not None else ""
    return {k: v for k, v in said.items() if v}


@dataclass
class Mail:
    id: str
    thread: str
    sender: str
    to: str
    subject: str
    body: str
    when: _dt.datetime
    labels: set[str] = field(default_factory=set)
    cc: str = ""
    message_id: str = ""
    in_reply_to: str = ""
    bcc: str = ""  # kept on the sender's own copy, as Gmail keeps it


class Gmail:
    """The user's Gmail, in memory, answering the API's calls.

    With `ignores_filters` a search answers every message, whatever it
    asked, as Google's CalDAV ignored a UID filter on 11 October 2026: what
    a provider keeps of a search is then only what it checked itself."""

    def __init__(self, world: World, *, ignores_filters: bool = False) -> None:
        self.world = world
        self.messages: dict[str, Mail] = {}
        self.drafts: dict[str, str] = {}  # draft id -> message id
        self.ignores_filters = ignores_filters
        self._ids = itertools.count(1)

    def new_id(self) -> str:
        return f"19a{next(self._ids):013x}"

    def put(self, *, sender: str, to: str, subject: str, body: str, when: _dt.datetime,
            labels: set[str], thread: str | None = None, message_id: str = "") -> Mail:
        mail = Mail(id=self.new_id(), thread="", sender=sender, to=to, subject=subject,
                    body=body, when=when, labels=set(labels), message_id=message_id)
        mail.thread = thread or mail.id
        self.messages[mail.id] = mail
        return mail

    # -- the search language, as much of it as Steve speaks -----------------

    def _operator(self, mail: Mail, op: str, value: str) -> bool:
        text = value.lower()
        if op == "is":
            return {"unread": "UNREAD" in mail.labels, "read": "UNREAD" not in mail.labels,
                    "starred": "STARRED" in mail.labels}.get(text, False)
        if op == "in":
            return text == "anywhere" or text.upper() in mail.labels
        if op == "label":
            return text.upper() in mail.labels
        if op == "category":
            return text == "primary"  # no message here is filed in another
        if op in ("newer_than", "older_than"):
            amount, unit = int(text[:-1]), text[-1]
            delta = _dt.timedelta(days=amount * {"d": 1, "m": 31, "y": 366}[unit])
            newer = mail.when >= MOMENT - delta
            return newer if op == "newer_than" else not newer
        if op in ("after", "before") and text.isdigit():
            # Seconds since the epoch, which Gmail takes as well as a day,
            # and which the provider writes for a search's `since`.
            moment = _dt.datetime.fromtimestamp(int(text), UTC)
            return mail.when >= moment if op == "after" else mail.when < moment
        if op in ("after", "before"):
            day = _dt.date(*map(int, text.split("/")))
            mine = mail.when.astimezone(ZONE).date()
            return mine >= day if op == "after" else mine < day
        if op in ("from", "to", "cc", "subject"):
            have = {"from": mail.sender, "to": mail.to, "cc": mail.cc,
                    "subject": mail.subject}[op]
            return text in have.lower()
        if op == "rfc822msgid":
            return text.strip("<>") == mail.message_id.lower().strip("<>")
        self.world.unhandled.append(f"gmail search operator {op}:")
        return True

    def _term(self, mail: Mail, term: str) -> bool:
        negate = term.startswith("-") and len(term) > 1
        term = term[1:] if negate else term
        if found := re.fullmatch(r"(\w+):(.+)", term):
            hit = self._operator(mail, found.group(1).lower(), found.group(2))
        else:
            words = f"{mail.subject} {mail.body} {mail.sender} {mail.to}".lower()
            hit = term.lower() in words
        return hit != negate

    def matches(self, mail: Mail, query: str) -> bool:
        try:
            terms = shlex.split(query or "")
        except ValueError:
            terms = (query or "").split()
        scoped = any(re.match(r"-?in:", t) for t in terms)
        if not scoped and mail.labels & {"DRAFT", "TRASH", "SPAM"}:
            return False
        hits: list[bool] = []
        either = False
        for term in terms:
            if term == "OR":
                either = True
                continue
            hit = self._term(mail, term)
            if either and hits:
                hits[-1] = hits[-1] or hit
            else:
                hits.append(hit)
            either = False
        return all(hits)

    # -- the API ---------------------------------------------------------------

    def resource(self, mail: Mail, fmt: str = "full", wanted=None) -> dict:
        headers = [("From", mail.sender), ("To", mail.to), ("Subject", mail.subject),
                   ("Date", email.utils.format_datetime(mail.when)),
                   ("Message-ID", mail.message_id)]
        if mail.cc:
            headers.append(("Cc", mail.cc))
        if mail.bcc:
            headers.append(("Bcc", mail.bcc))
        if mail.in_reply_to:
            headers += [("In-Reply-To", mail.in_reply_to), ("References", mail.in_reply_to)]
        shown = {"id": mail.id, "threadId": mail.thread, "labelIds": sorted(mail.labels),
                 "snippet": mail.body[:100], "sizeEstimate": len(mail.body),
                 "internalDate": str(int(mail.when.timestamp() * 1000))}
        if fmt == "minimal":
            return shown
        if wanted:
            names = {w.lower() for w in wanted}
            headers = [h for h in headers if h[0].lower() in names]
        payload: dict = {"mimeType": "text/plain",
                         "headers": [{"name": n, "value": v} for n, v in headers if v]}
        if fmt != "metadata":
            payload["body"] = {"size": len(mail.body), "data": _b64(mail.body)}
        shown["payload"] = payload
        return shown

    def profile(self, userId):  # noqa: N803 - the API's names
        return {"emailAddress": ME, "messagesTotal": len(self.messages),
                "threadsTotal": len({m.thread for m in self.messages.values()})}

    def list(self, userId, q="", maxResults=100, pageToken=None):  # noqa: N803
        found = sorted((m for m in self.messages.values()
                        if self.ignores_filters or self.matches(m, q)),
                       key=lambda m: m.when, reverse=True)[: int(maxResults)]
        return {"messages": [{"id": m.id, "threadId": m.thread} for m in found],
                "resultSizeEstimate": len(found)}

    def get(self, userId, id, format="full", metadataHeaders=None):  # noqa: A002, N803
        if id not in self.messages:
            raise _not_found(f"gmail message {id}")
        if format == "raw":
            return self.raw(self.messages[id])
        return self.resource(self.messages[id], format, metadataHeaders)

    def raw(self, mail: Mail) -> dict:
        """A message whole, as `format="raw"` gives it: its RFC 822 bytes."""
        message = email.message.EmailMessage()
        for name, value in (("From", mail.sender), ("To", mail.to), ("Cc", mail.cc),
                            ("Bcc", mail.bcc), ("Subject", mail.subject),
                            ("Date", email.utils.format_datetime(mail.when)),
                            ("Message-ID", mail.message_id), ("In-Reply-To", mail.in_reply_to),
                            ("References", mail.in_reply_to)):
            if value:
                message[name] = value
        message.set_content(mail.body)
        return {"id": mail.id, "threadId": mail.thread, "labelIds": sorted(mail.labels),
                "raw": _b64(message.as_bytes())}

    def save_draft_again(self, draft_id: str, **changes) -> Mail:
        """A draft saved again, as from the user's browser: Gmail never
        changes a message, so the draft holds a new one, under a new id,
        and the one it held is gone."""
        old = self.messages.pop(self.drafts[draft_id])
        new = dataclasses.replace(old, id=self.new_id(), labels=set(old.labels), **changes)
        self.messages[new.id] = new
        self.drafts[draft_id] = new.id
        return new

    def thread(self, userId, id, format="full", metadataHeaders=None):  # noqa: A002, N803
        found = sorted((m for m in self.messages.values() if m.thread == id),
                       key=lambda m: m.when)
        if not found:
            raise _not_found(f"gmail thread {id}")
        return {"id": id, "messages": [self.resource(m, format, metadataHeaders) for m in found]}

    def modify(self, userId, id, body):  # noqa: A002, N803
        mail = self.messages[id]
        mail.labels |= set(body.get("addLabelIds") or [])
        mail.labels -= set(body.get("removeLabelIds") or [])
        self.world.wrote.append({"gmail": "messages.modify", "id": id, **body})
        return self.resource(mail, "minimal")

    def _from_raw(self, message: dict, labels: set[str]) -> Mail:
        said = _written(message["raw"])
        mail = self.put(sender=said.get("from", ME), to=said.get("to", ""),
                        subject=said.get("subject", ""), body=said.get("body", ""),
                        when=MOMENT, labels=labels, thread=message.get("threadId"))
        mail.cc, mail.in_reply_to = said.get("cc", ""), said.get("in_reply_to", "")
        # What `_written` leaves out of the record: the Bcc, which Gmail keeps
        # on the sender's copy, and the Message-ID, which Gmail gives a
        # message written without one.
        raw = email.message_from_bytes(_unb64(message["raw"]), policy=email.policy.default)
        mail.bcc = str(raw.get("Bcc") or "")
        mail.message_id = str(raw.get("Message-ID") or f"<{mail.id}.golden@mail.example>")
        return mail

    def send(self, userId, body):  # noqa: N803
        mail = self._from_raw(body, {"SENT"})
        self.world.wrote.append({"gmail": "messages.send", "thread": body.get("threadId"),
                                 **_written(body["raw"])})
        return {"id": mail.id, "threadId": mail.thread, "labelIds": ["SENT"]}

    def create_draft(self, userId, body):  # noqa: N803
        mail = self._from_raw(body["message"], {"DRAFT"})
        draft = f"r{next(self._ids):019d}"
        self.drafts[draft] = mail.id
        self.world.wrote.append({"gmail": "drafts.create", "draft": draft,
                                 "thread": body["message"].get("threadId"),
                                 **_written(body["message"]["raw"])})
        return {"id": draft, "message": {"id": mail.id, "threadId": mail.thread,
                                         "labelIds": ["DRAFT"]}}

    def get_draft(self, userId, id, format="full"):  # noqa: A002, N803
        if id not in self.drafts:
            raise _not_found(f"gmail draft {id}")
        return {"id": id, "message": self.resource(self.messages[self.drafts[id]], format)}

    def send_draft(self, userId, body):  # noqa: N803
        if body["id"] not in self.drafts:
            raise _not_found(f"gmail draft {body['id']}")
        mail = self.messages[self.drafts.pop(body["id"])]
        mail.labels = {"SENT"}
        self.world.wrote.append({"gmail": "drafts.send", "draft": body["id"]})
        return {"id": mail.id, "threadId": mail.thread, "labelIds": ["SENT"]}

    def delete_draft(self, userId, id):  # noqa: A002, N803
        if id not in self.drafts:
            raise _not_found(f"gmail draft {id}")
        del self.messages[self.drafts.pop(id)]
        self.world.wrote.append({"gmail": "drafts.delete", "draft": id})
        return ""


def _start(event: dict, edge: str = "start") -> _dt.datetime:
    side = event[edge]
    if "dateTime" in side:
        return _dt.datetime.fromisoformat(side["dateTime"])
    return _dt.datetime.combine(_dt.date.fromisoformat(side["date"]), _dt.time(), ZONE)


def _rfc3339(moment: _dt.datetime) -> str:
    return moment.astimezone(UTC).isoformat().replace("+00:00", "Z")


class Calendar:
    """The user's Google Calendar: one calendar, `primary`."""

    def __init__(self, world: World) -> None:
        self.world = world
        self.events: list[dict] = []
        self._etags = itertools.count(1)

    def put(self, summary: str, start: _dt.datetime, minutes: int = 60, *,
            guests: tuple[str, ...] = (), all_day: bool = False,
            updated: _dt.datetime | None = None,
            **extra) -> dict:
        eid = f"ev{len(self.events) + 1:04d}golden"
        if all_day:
            span = {"start": {"date": start.date().isoformat()},
                    "end": {"date": (start.date() + _dt.timedelta(days=1)).isoformat()}}
        else:
            end = start + _dt.timedelta(minutes=minutes)
            span = {"start": {"dateTime": start.isoformat(), "timeZone": persona.ZONE},
                    "end": {"dateTime": end.isoformat(), "timeZone": persona.ZONE}}
        attendees = [{"email": ME, "self": True, "responseStatus": "accepted"}]
        attendees += [{"email": g, "responseStatus": "accepted"} for g in guests]
        # Google's own form is the id at google.com; a reserved domain here,
        # since a card binds the UID and the record would then hold it.
        event = {"id": eid, "iCalUID": f"{eid}@google.example",
                 "etag": f'"3{len(self.events):05d}"',
                 "status": "confirmed", "summary": summary, **span,
                 "htmlLink": f"https://calendar.example/event?eid={eid}",
                 "created": _rfc3339(at(days=-20)),
                 "updated": _rfc3339(updated or at(days=-10)),
                 "organizer": {"email": ME, "self": True},
                 "attendees": attendees if guests else [], **extra}
        self.events.append(event)
        return event

    def setting(self, setting):
        return {"kind": "calendar#setting", "id": setting, "value": persona.ZONE}

    def list(self, calendarId, singleEvents=True, timeMin=None, timeMax=None,  # noqa: N803
             orderBy=None, maxResults=250, q=None, updatedMin=None, showDeleted=False,
             timeZone=None):
        low = _dt.datetime.fromisoformat(timeMin) if timeMin else None
        high = _dt.datetime.fromisoformat(timeMax) if timeMax else None
        since = _dt.datetime.fromisoformat(updatedMin) if updatedMin else None
        found = []
        listed = [one for event in self.events
                  for one in (self._occurrences(event) if singleEvents else [event])]
        for event in listed:
            if event["status"] == "cancelled" and not (showDeleted or updatedMin):
                continue
            if low and _start(event, "end") <= low or high and _start(event) >= high:
                continue
            if since and _dt.datetime.fromisoformat(event["updated"]) < since:
                continue
            if q and q.lower() not in " ".join(
                    str(event.get(k) or "") for k in ("summary", "description", "location")
            ).lower():
                continue
            found.append(event)
        found.sort(key=_start)
        return {"items": found[: int(maxResults)], "timeZone": persona.ZONE}

    _RULE = re.compile(r"RRULE:FREQ=(DAILY|WEEKLY);COUNT=(\d+)")

    def _occurrences(self, event: dict) -> list[dict]:
        """A repeating event as Google lists it with `singleEvents`: each
        occurrence its own id (the series' and its start), the series' UID,
        and `recurringEventId`. Only the daily and weekly counts are known."""
        rules = event.get("recurrence") or []
        if not rules:
            return [event]
        found = self._RULE.fullmatch(rules[0])
        if found is None:
            self.world.unhandled.append(f"calendar recurrence {rules[0]}")
            return [event]
        step = _dt.timedelta(days=1 if found.group(1) == "DAILY" else 7)
        start, end = _start(event), _start(event, "end")
        out = []
        for n in range(int(found.group(2))):
            at_ = start + n * step
            one = {k: v for k, v in event.items() if k != "recurrence"}
            one.update(id=f"{event['id']}_{at_.astimezone(UTC):%Y%m%dT%H%M%SZ}",
                       recurringEventId=event["id"],
                       originalStartTime={"dateTime": at_.isoformat()},
                       start={**event["start"], "dateTime": at_.isoformat()},
                       end={**event["end"], "dateTime": (end + n * step).isoformat()})
            out.append(one)
        return out

    def get(self, calendarId, eventId):  # noqa: N803
        for event in self.events:
            if event["id"] == eventId:
                return event
            if eventId.startswith(event["id"] + "_"):
                for one in self._occurrences(event):
                    if one["id"] == eventId:
                        return one
        raise _not_found(f"calendar event {eventId}")

    def _held(self, event_id: str, if_match: str | None) -> dict:
        """The event a write names, refused as Google refuses it (412) when
        it is conditional on an etag the event no longer has."""
        event = self.get("primary", event_id)
        if if_match is not None and if_match != event["etag"]:
            raise _precondition_failed(f"calendar event {event_id}")
        return event

    def _touched(self, event: dict) -> None:
        event["etag"] = f'"4{next(self._etags):05d}"'
        event["updated"] = _rfc3339(MOMENT)

    def insert(self, calendarId, body, sendUpdates="none"):  # noqa: N803
        start = _dt.datetime.fromisoformat(body["start"]["dateTime"])
        minutes = int((_dt.datetime.fromisoformat(body["end"]["dateTime"]) - start)
                      .total_seconds() // 60)
        event = self.put(body["summary"], start, minutes,
                         guests=tuple(a["email"] for a in body.get("attendees") or ()),
                         **{k: body[k] for k in ("location", "description", "visibility")
                            if k in body})
        self._touched(event)
        self.world.wrote.append({"calendar": "events.insert", "sendUpdates": sendUpdates,
                                 **body})
        return event

    def patch(self, calendarId, eventId, body, if_match=None):  # noqa: N803
        event = self._held(eventId, if_match)
        event.update(body)
        self._touched(event)
        self.world.wrote.append({"calendar": "events.patch", "id": eventId,
                                 "if_match": if_match, **body})
        return event

    def delete(self, calendarId, eventId, sendUpdates="none", if_match=None):  # noqa: N803
        event = self._held(eventId, if_match)
        event["status"] = "cancelled"
        self._touched(event)
        self.world.wrote.append({"calendar": "events.delete", "id": eventId,
                                 "sendUpdates": sendUpdates, "if_match": if_match})
        return ""

    def freebusy(self, body):
        low = _dt.datetime.fromisoformat(body["timeMin"])
        high = _dt.datetime.fromisoformat(body["timeMax"])
        busy = []
        for event in sorted(self.events, key=_start):
            if event["status"] == "cancelled" or "date" in event["start"]:
                continue
            start, end = max(_start(event), low), min(_start(event, "end"), high)
            if start < end:
                busy.append({"start": _rfc3339(start), "end": _rfc3339(end)})
        return {"calendars": {"primary": {"busy": busy}}}


class Drive:
    def __init__(self, world: World) -> None:
        self.world = world
        self.files: list[dict] = []

    def put(self, name: str, mime: str, owner: str, modified: _dt.datetime, text: str) -> dict:
        fid = f"1golden{len(self.files) + 1:04d}Drive"
        entry = {"id": fid, "name": name, "mimeType": mime, "modifiedTime": _rfc3339(modified),
                 "webViewLink": f"https://docs.example/d/{fid}",
                 "owners": [{"emailAddress": owner}], "_text": text}
        self.files.append(entry)
        return entry

    _CONTAINS = re.compile(r"(?:name|fullText) contains '((?:\\.|[^'\\])*)'")

    def list(self, q="", pageSize=100, fields=None, orderBy=None, pageToken=None):  # noqa: N803
        rest = re.sub(r"\(|\)|\band\b|\bor\b|trashed = false", " ", self._CONTAINS.sub(" ", q))
        mime = re.search(r"mimeType = '([^']*)'", rest)
        rest = re.sub(r"mimeType = '[^']*'", " ", rest)
        if rest.strip():
            self.world.unhandled.append(f"drive query {rest.strip()!r}")
        words = [re.sub(r"\\(.)", r"\1", w).lower() for w in self._CONTAINS.findall(q)]
        found = [f for f in self.files
                 if all(w in (f["name"] + " " + f["_text"]).lower() for w in words)
                 and (mime is None or f["mimeType"] == mime.group(1))]
        found.sort(key=lambda f: f["modifiedTime"], reverse=True)
        return {"files": [self._shown(f) for f in found[: int(pageSize)]]}

    @staticmethod
    def _shown(entry: dict) -> dict:
        return {k: v for k, v in entry.items() if not k.startswith("_")}

    def _file(self, fid: str) -> dict:
        for entry in self.files:
            if entry["id"] == fid:
                return entry
        raise _not_found(f"drive file {fid}")

    def get(self, fileId, fields=None):  # noqa: N803
        return self._shown(self._file(fileId))

    def export(self, fileId, mimeType):  # noqa: N803
        return self._file(fileId)["_text"].encode()

    def create(self, body, media_body=None, fields=None):
        """A file uploaded and converted, as Docs are made: a .docx's
        paragraphs become the Google Doc's text."""
        text = ""
        if media_body is not None:
            import docx

            data = media_body.getbytes(0, media_body.size())
            text = "\n\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        entry = self.put(body["name"], body.get("mimeType", ""), ME, MOMENT, text)
        self.world.wrote.append({"drive": "files.create", "name": body["name"]})
        return self._shown(entry)

    # -- Sheets, whose spreadsheets are Drive's files -------------------------

    @staticmethod
    def _cell(ref: str) -> tuple[int, int]:
        letters, digits = re.fullmatch(r"([A-Z]+)(\d+)", ref).groups()
        col = 0
        for letter in letters:
            col = col * 26 + ord(letter) - ord("A") + 1
        return int(digits) - 1, col - 1

    def _range(self, a1: str) -> tuple[tuple[int, int], tuple[int, int] | None]:
        cells = a1.rpartition("!")[2]  # a sheet's name, or the first sheet
        first, _, last = cells.partition(":")
        return self._cell(first), self._cell(last) if last else None

    def create_sheet(self, body, fields=None):
        entry = self.put(body["properties"]["title"], "application/vnd.google-apps.spreadsheet",
                         ME, MOMENT, "")
        entry["_cells"] = {}
        self.world.wrote.append({"sheets": "spreadsheets.create",
                                 "title": body["properties"]["title"]})
        return {"spreadsheetId": entry["id"], "spreadsheetUrl": entry["webViewLink"]}

    def values_get(self, spreadsheetId, range):  # noqa: A002, N803 - the API's names
        return {"range": range, "values": self._values(self._file(spreadsheetId)["_cells"],
                                                       range)}

    def _values(self, cells: dict, a1: str) -> list[list[str]]:
        """The cells asked for, as Sheets answers: no empty cells at a row's
        end, and no empty rows at the range's."""
        (top, left), end = self._range(a1)
        bottom, right = end or (top, left)
        rows = []
        for r in range(top, bottom + 1):
            row = [cells.get((r, c), "") for c in range(left, right + 1)]
            while row and row[-1] == "":
                row.pop()
            rows.append(row)
        while rows and not rows[-1]:
            rows.pop()
        return rows

    def values_update(self, spreadsheetId, range, valueInputOption, body):  # noqa: A002, N803
        entry = self._file(spreadsheetId)
        (top, left), _ = self._range(range)
        for r, row in enumerate(body["values"]):
            for c, value in enumerate(row):
                entry["_cells"][(top + r, left + c)] = str(value)
        entry["modifiedTime"] = _rfc3339(MOMENT)
        written = sum(len(row) for row in body["values"])
        self.world.wrote.append({"sheets": "values.update", "id": spreadsheetId,
                                 "range": range, "values": body["values"]})
        return {"updatedCells": written, "updatedRange": range}


class Tasks:
    def __init__(self, world: World) -> None:
        self.world = world
        self.lists: list[dict] = []
        self.items: dict[str, list[dict]] = {}
        self._ids = itertools.count(1)

    def put_list(self, title: str) -> str:
        lid = f"MTgolden{len(self.lists) + 1:04d}"
        self.lists.append({"id": lid, "title": title})
        self.items[lid] = []
        return lid

    def put(self, lid: str, title: str, *, due: str | None = None, done: bool = False) -> dict:
        task = {"id": f"dGFza{next(self._ids):04d}golden", "title": title,
                "status": "completed" if done else "needsAction"}
        if due:
            task["due"] = f"{due}T00:00:00.000Z"
        self.items[lid].append(task)
        return task

    def tasklists(self, maxResults=100, pageToken=None):  # noqa: N803
        return {"items": list(self.lists)}

    def insert_list(self, body):
        lid = self.put_list(body["title"])
        self.world.wrote.append({"tasks": "tasklists.insert", "title": body["title"]})
        return {"id": lid, "title": body["title"]}

    def list(self, tasklist, showCompleted=True, showHidden=False, maxResults=100,  # noqa: N803
             pageToken=None):
        found = [t for t in self.items[tasklist]
                 if showCompleted or t["status"] != "completed"]
        return {"items": found[: int(maxResults)]}

    def _task(self, tasklist: str, task: str) -> dict:
        for item in self.items.get(tasklist, []):
            if item["id"] == task:
                return item
        raise _not_found(f"task {task}")

    def get(self, tasklist, task):
        return self._task(tasklist, task)

    def insert(self, tasklist, body):
        task = self.put(tasklist, body["title"], due=(body.get("due") or "")[:10] or None)
        if body.get("notes"):
            task["notes"] = body["notes"]
        self.world.wrote.append({"tasks": "tasks.insert", "list": tasklist, **body})
        return task

    def patch(self, tasklist, task, body):
        found = self._task(tasklist, task)
        found.update(body)
        self.world.wrote.append({"tasks": "tasks.patch", "list": tasklist, "task": task, **body})
        return found


class Google:
    """`googleapiclient.discovery.build`, answering from the world."""

    def __init__(self, world: World) -> None:
        self.world = world
        g, c, d, t = world.gmail, world.calendar, world.drive, world.tasks
        self.handlers = {
            "gmail.users.getProfile": g.profile,
            "gmail.users.messages.list": g.list,
            "gmail.users.messages.get": g.get,
            "gmail.users.messages.send": g.send,
            "gmail.users.messages.modify": g.modify,
            "gmail.users.threads.get": g.thread,
            "gmail.users.drafts.create": g.create_draft,
            "gmail.users.drafts.get": g.get_draft,
            "gmail.users.drafts.send": g.send_draft,
            "gmail.users.drafts.delete": g.delete_draft,
            "calendar.settings.get": c.setting,
            "calendar.events.list": c.list,
            "calendar.events.get": c.get,
            "calendar.events.insert": c.insert,
            "calendar.events.patch": c.patch,
            "calendar.events.delete": c.delete,
            "calendar.freebusy.query": c.freebusy,
            "drive.files.list": d.list,
            "drive.files.get": d.get,
            "drive.files.export": d.export,
            "drive.files.create": d.create,
            "sheets.spreadsheets.create": d.create_sheet,
            "sheets.spreadsheets.values.get": d.values_get,
            "sheets.spreadsheets.values.update": d.values_update,
            "tasks.tasklists.list": t.tasklists,
            "tasks.tasklists.insert": t.insert_list,
            "tasks.tasks.list": t.list,
            "tasks.tasks.get": t.get,
            "tasks.tasks.insert": t.insert,
            "tasks.tasks.patch": t.patch,
        }

    def build(self, api: str, version: str, **_options) -> _Resource:
        return _Resource(self, api)

    def unhandled(self, what: str) -> Exception:
        self.world.unhandled.append(what)
        return LookupError(f"the golden world has no {what}; add it to {__name__}")


class _Resource:
    """One level of a client: `users()`, `messages()`, then a method."""

    def __init__(self, google: Google, path: str) -> None:
        self._google, self._path = google, path

    def __getattr__(self, name: str):
        if name.startswith("_"):
            raise AttributeError(name)
        path = f"{self._path}.{name}"
        if (handler := self._google.handlers.get(path)) is not None:
            return lambda **kwargs: _Request(self._google, path, handler, kwargs)
        if any(key.startswith(path + ".") for key in self._google.handlers):
            return lambda: _Resource(self._google, path)
        raise self._google.unhandled(path)


class _Request:
    """A request as the client builds it, with the headers a caller may set
    on it before it runs; an If-Match reaches its handler as `if_match`."""

    def __init__(self, google: Google, path: str, handler, kwargs: dict) -> None:
        self._google, self._path, self._handler, self._kwargs = google, path, handler, kwargs
        self.headers: dict[str, str] = {}

    def execute(self, http=None, num_retries=0):
        kwargs = dict(self._kwargs)
        if "If-Match" in self.headers:
            kwargs["if_match"] = self.headers["If-Match"]
        try:
            bound = inspect.signature(self._handler).bind(**kwargs)
        except TypeError:
            raise self._google.unhandled(
                f"{self._path}({', '.join(sorted(self._kwargs))})") from None
        return self._handler(*bound.args, **bound.kwargs)


# ---------------------------------------------------------------------------
# Outlook, through MAPI
# ---------------------------------------------------------------------------


class _Collection:
    def __init__(self, items: list) -> None:
        self._items = items
        self.Count = len(items)

    def Item(self, i):  # noqa: N802 - COM's names, here and below
        return self._items[i - 1]


class _Recipient:
    def __init__(self, name: str, address: str, kind: int = 1) -> None:
        self.Name, self.Address, self.Type = name, address, kind
        self.MeetingResponseStatus = 3

    @property
    def AddressEntry(self):  # noqa: N802
        return types.SimpleNamespace(
            Address=self.Address,
            GetExchangeUser=lambda: types.SimpleNamespace(PrimarySmtpAddress=self.Address))


class OutlookMail:
    """A MailItem: a message in a folder, or a draft being written."""

    def __init__(self, outlook: Outlook, *, subject: str = "", sender=("", ""), to=(),
                 body: str = "", when: _dt.datetime | None = None, unread: bool = False,
                 conversation: str = "") -> None:
        self._outlook = outlook
        self.Subject, self.Body, self.UnRead, self.Class = subject, body, unread, 43
        # MAPI's mail times are the local wall clock, naive
        # (modules/outlook_classic/outlook._epoch_ms).
        wall = (when or MOMENT).astimezone(ZONE).replace(tzinfo=None)
        self.SentOn = self.ReceivedTime = wall
        self.SenderName, self.SenderEmailAddress = sender
        # Every recipient, each of its type: To, Cc or Bcc.
        self._to = [_Recipient(n, a) for n, a in to]
        self._lines: dict[int, str] = {}
        self.HTMLBody = f"<html><body><p>{body}</p></body></html>" if body else ""
        self.EntryID = ""
        # Not Outlook's: the id it was first filed under, which a world keys
        # it by once Exchange has filed it again under another.
        self.made_as = ""
        self.ConversationID = conversation
        self.Sent = False
        self.SendUsingAccount = None

    @property
    def Sender(self):  # noqa: N802
        if not self.SenderEmailAddress:
            return None
        return _Recipient(self.SenderName, self.SenderEmailAddress)

    @property
    def Recipients(self):  # noqa: N802
        return _Collection(self._to)

    def _line(self, kind: int) -> str:
        if kind in self._lines:
            return self._lines[kind]
        return "; ".join(r.Name for r in self._to if r.Type == kind)

    def _resolve(self, kind: int, value) -> None:
        # Outlook resolves what is written to `.To`, `.CC` or `.BCC` into the
        # Recipients collection, as recipients of that type, which is where
        # a fetch reads the addresses from.
        self._lines[kind] = value
        self._to = [r for r in self._to if r.Type != kind] + [
            _Recipient(name, address, kind)
            for name, address in email.utils.getaddresses([value or ""]) if address]

    @property
    def To(self):  # noqa: N802
        return self._line(1)

    @To.setter
    def To(self, value):  # noqa: N802
        self._resolve(1, value)

    @property
    def CC(self):  # noqa: N802
        return self._line(2)

    @CC.setter
    def CC(self, value):  # noqa: N802
        self._resolve(2, value)

    @property
    def BCC(self):  # noqa: N802
        return self._line(3)

    @BCC.setter
    def BCC(self, value):  # noqa: N802
        self._resolve(3, value)

    def _reply(self) -> OutlookMail:
        reply = OutlookMail(self._outlook, subject=f"RE: {self.Subject}",
                            to=[(self.SenderName, self.SenderEmailAddress)],
                            conversation=self.ConversationID)
        reply.HTMLBody = (f"<html><body><div><p>From: {self.SenderName}</p>"
                          f"<p>{self.Body}</p></div></body></html>")
        return reply

    def Reply(self):  # noqa: N802
        return self._reply()

    def ReplyAll(self):  # noqa: N802
        return self._reply()

    def Save(self):  # noqa: N802
        self._outlook.save(self)

    def Send(self):  # noqa: N802
        self._outlook.send(self)

    def Delete(self):  # noqa: N802
        self._outlook.delete(self)


class Appointment:
    def __init__(self, subject: str, start: _dt.datetime, minutes: int = 60, *,
                 guests=(), location: str = "", body: str = "") -> None:
        self.Subject = subject
        self.StartUTC = start.astimezone(UTC)
        self.EndUTC = (start + _dt.timedelta(minutes=minutes)).astimezone(UTC)
        # `Start` is the wall clock, what the folder sorts by; the reader takes StartUTC.
        self.Start = start.astimezone(ZONE).replace(tzinfo=None)
        self.AllDayEvent, self.MeetingStatus, self.IsRecurring = False, 1, False
        self.EntryID = f"00000000APPT{subject.replace(' ', '')}"
        self.GlobalAppointmentID = f"040000008200E000{subject.replace(' ', '')}"
        self.Recipients = _Collection([_Recipient(n, a) for n, a in guests])
        self.Location, self.Body = location, body
        self.LastModificationTime = self.CreationTime = (
            at(days=-5).astimezone(ZONE).replace(tzinfo=None))

    def GetOrganizer(self):  # noqa: N802
        return types.SimpleNamespace(Address=persona.WORK_EMAIL)

    def repeat(self, every: _dt.timedelta, count: int) -> None:
        """Make this a series of `count` occurrences, `every` apart."""
        self.IsRecurring, self._every, self._count = True, every, count

    def occurrences(self) -> list[Appointment]:
        """The series as a view with `IncludeRecurrences` lists it: each
        occurrence the series' own ids at its own start."""
        if not self.IsRecurring:
            return [self]
        out = []
        for n in range(self._count):
            one, shift = copy.copy(self), n * self._every
            one.StartUTC, one.EndUTC = self.StartUTC + shift, self.EndUTC + shift
            one.Start = one.StartUTC.astimezone(ZONE).replace(tzinfo=None)
            out.append(one)
        return out


class _Items:
    """A folder's Items: a fresh view each time it is asked for, as MAPI's is.
    Sorted and expanded when walked, so `IncludeRecurrences` set after `Sort`
    still counts, as Outlook's own order of the two allows."""

    def __init__(self, items: list) -> None:
        self._all, self._view, self._at = items, list(items), 0
        self._key: str | None = None
        self._descending = False
        self.IncludeRecurrences = False

    Count = property(lambda self: len(self._listed()))

    def _listed(self) -> list:
        items = self._all
        if self.IncludeRecurrences:
            items = [one for item in items
                     for one in (item.occurrences() if hasattr(item, "occurrences") else [item])]
        if self._key is None:
            return list(items)
        return sorted(items, key=lambda i: getattr(i, self._key), reverse=self._descending)

    def Sort(self, key, descending=False):  # noqa: N802
        self._key, self._descending = key.strip("[]"), bool(descending)
        self._at = 0

    def GetFirst(self):  # noqa: N802
        self._view, self._at = self._listed(), 0
        return self.GetNext()

    def GetNext(self):  # noqa: N802
        if self._at >= len(self._view):
            return None
        self._at += 1
        return self._view[self._at - 1]


class _Folder:
    def __init__(self, name: str) -> None:
        self.Name, self.EntryID, self.items = name, f"0000FOLDER{name.replace(' ', '')}", []
        self.Folders = _Collection([])

    @property
    def Items(self):  # noqa: N802
        return _Items(self.items)

    @property
    def UnReadItemCount(self):  # noqa: N802
        return sum(1 for i in self.items if getattr(i, "UnRead", False))


class Outlook:
    """Outlook, through MAPI: one store, one account, the user's work mail."""

    INBOX, SENT, DRAFTS, CALENDAR = 6, 5, 16, 9

    def __init__(self, world: World) -> None:
        self.world = world
        self.folders = {fid: _Folder(name) for fid, name in (
            (self.INBOX, "Inbox"), (self.SENT, "Sent Items"), (self.DRAFTS, "Drafts"),
            (self.CALENDAR, "Calendar"))}
        self._ids = itertools.count(1)
        account = types.SimpleNamespace(SmtpAddress=persona.WORK_EMAIL)
        self.Accounts = _Collection([account])
        store = types.SimpleNamespace(DisplayName=persona.WORK_EMAIL,
                                      GetDefaultFolder=self.GetDefaultFolder)
        self.Stores = _Collection([store])
        self.Application = types.SimpleNamespace(CreateItem=self._create)

    def put(self, folder: int, item) -> None:
        if isinstance(item, OutlookMail) and not item.EntryID:
            item.EntryID = item.made_as = self._entry_id()
        self.folders[folder].items.append(item)

    def _entry_id(self) -> str:
        # Hex, as MAPI's are: the normaliser knows an Outlook id by it.
        return f"00000000A1B2C3D4E5F60708090A0B0C0D0E0F{next(self._ids):08X}"

    def _create(self, kind: int) -> OutlookMail:
        assert kind == 0, "only mail items are made here"
        return OutlookMail(self)

    def GetDefaultFolder(self, fid: int):  # noqa: N802
        return self.folders[fid]

    def GetItemFromID(self, entry_id: str):  # noqa: N802
        for folder in self.folders.values():
            for item in folder.items:
                if item.EntryID == entry_id:
                    return item
        raise RuntimeError("(-2147221233, 'The operation failed.')")

    def _where(self, item) -> _Folder | None:
        return next((f for f in self.folders.values() if item in f.items), None)

    def save(self, item: OutlookMail) -> None:
        if self._where(item) is None:
            self.put(self.DRAFTS, item)
        self.world.wrote.append({"outlook": "Save", "to": item.To, "subject": item.Subject,
                                 "html": item.HTMLBody})

    def send(self, item: OutlookMail) -> None:
        if (folder := self._where(item)) is not None:
            folder.items.remove(item)
        item.Sent = True
        # Filed in Sent Items under an id of its own, as Outlook files the
        # copy it keeps: the draft's id finds nothing once it is sent.
        item.EntryID = self._entry_id()
        self.folders[self.SENT].items.append(item)
        self.world.wrote.append({"outlook": "Send", "to": item.To, "subject": item.Subject})

    def delete(self, item: OutlookMail) -> None:
        self._where(item).items.remove(item)
        self.world.wrote.append({"outlook": "Delete", "subject": item.Subject})


# ---------------------------------------------------------------------------
# The world: the user's accounts, filled in
# ---------------------------------------------------------------------------

# Their course site's calendar link, marked as due dates (step 5, item 8:
# it was a feed of its own, read by a tool of its own).
COURSE_LINK = "https://courses.example/feeds/calendars/user_x.ics?token=golden"
COURSE_FEED = """BEGIN:VCALENDAR
VERSION:2.0
PRODID:icalendar-ruby
X-WR-CALNAME:Maya Lindqvist Calendar
BEGIN:VEVENT
DTSTAMP:20261001T120000Z
UID:event-assignment-700101
DTSTART:20261017T035900Z
DTEND:20261017T035900Z
DESCRIPTION:Glaze three test tiles and bring them to the studio.
SUMMARY:Test tiles [ART 110]
URL;VALUE=URI:https://courses.example/calendar#assignment_700101
END:VEVENT
END:VCALENDAR
"""

IRIS = "Iris Novak <iris@fabrikam.example>"
DANA = "Dana Ortiz <dana@example.com>"
TOMAS = ("Tomas Varga", "tomas.varga@contoso.example")


class World:
    """Everything the user's accounts hold, and what was done to them."""

    def __init__(self) -> None:
        self.unhandled: list[str] = []
        self.wrote: list[dict] = []
        self.gmail, self.calendar = Gmail(self), Calendar(self)
        self.drive, self.tasks = Drive(self), Tasks(self)
        self.outlook = Outlook(self)
        # The user's own domain's mailbox: an account of `imap`, core's engine
        # against fake servers, over loopback sockets and TLS (step 5, item 6).
        self.mail = mailrig.rig(STANDARD, user=persona.OWN_DOMAIN_EMAIL)
        # Their course site's calendar link: an account of `ics`, core's
        # engine reading a feed served at its host (step 5, item 8).
        self.feeds = feeds.FeedRig()
        self.feeds.serve(COURSE_LINK, raw=COURSE_FEED)
        self.google = Google(self)
        self._fill()

    def _fill(self) -> None:
        g = self.gmail
        g.put(sender=IRIS, to=ME, subject="Quarterly plan draft",
              body="Could you look at the quarterly plan before Friday? "
                   "The numbers are in section 2.\n\nThanks,\nIris",
              when=at(hours=-1, minutes=-20), labels={"INBOX", "UNREAD"},
              message_id="<plan-1@fabrikam.example>")
        g.put(sender="Northwind Pharmacy <orders@northwind.example>", to=ME,
              subject="Your order has shipped",
              body="Your order 4471 has shipped and should arrive on Friday.",
              when=at(days=-1, hours=-2), labels={"INBOX"},
              message_id="<order-4471@northwind.example>")
        g.put(sender=ME, to=DANA, subject="Book club on Friday?",
              body="Are you still hosting book club on Friday? I can bring the second copy.",
              when=at(days=-6), labels={"SENT"}, message_id="<bookclub-1@gmail.example>")

        o = self.outlook
        o.put(o.INBOX, OutlookMail(o, subject="Badge renewal by Friday", sender=TOMAS,
                                   to=[("Maya Lindqvist", persona.WORK_EMAIL)],
                                   body="Security asked everyone to renew their building "
                                        "badge by Friday. The desk is open 9 to 4.",
                                   when=at(hours=-2), unread=True, conversation="CONV0001"))
        o.put(o.SENT, OutlookMail(o, subject="Vendor review notes", sender=(
            "Maya Lindqvist", persona.WORK_EMAIL), to=[TOMAS],
            body="Notes from the vendor review are attached to the ticket.",
            when=at(days=-3), conversation="CONV0002"))
        o.put(o.CALENDAR, Appointment("Design review with Iris", at(hours=1, minutes=30), 60,
                                      guests=[TOMAS], location="Room 4B",
                                      body="Walk through the quarterly plan."))
        o.put(o.CALENDAR, Appointment("Vendor call", at(days=1, hours=4, minutes=30), 30,
                                      guests=[TOMAS]))

        own = persona.OWN_DOMAIN_EMAIL
        self.mail.world.seed("mail", [
            {"sender": "Greenway Garden <committee@greenway-garden.example>", "to": own,
             "subject": "Plot assignments for spring",
             "body": "Your plot is number 14. The first work day is on Saturday at 10.",
             "date": at(hours=-3), "message_id": "<plots-14@greenway-garden.example>",
             "unread": True},
            {"role": "sent", "sender": own, "to": DANA, "subject": "The soup recipe",
             "body": "Here is the soup recipe I promised.", "date": at(days=-2),
             "message_id": "<soup-1@lindqvist.example>"}])
        # What reached the servers, in order, beside what reached the others.
        self.mail.world.on_write = self.wrote.append
        self.mail.smtp.on_send = self.wrote.append

        c = self.calendar
        c.put("Planning sync", at(hours=-2), 60, guests=("iris@fabrikam.example",
                                                        "tomas.varga@contoso.example"))
        c.put("Team standup", at(minutes=30), 15, guests=("iris@fabrikam.example",),
              updated=at(hours=-5))
        c.put("Design review with Iris", at(hours=1, minutes=30), 60,
              guests=("iris@fabrikam.example",), location="Room 4B")
        c.put("Lunch with Dana", at(hours=3), 60, guests=("dana@example.com",),
              location="Juniper Cafe", updated=at(hours=-20))
        c.put("Dentist", at(days=1, hours=5, minutes=30), 60, location="Harbor Dental")
        c.put("Choir rehearsal", at(hours=9), 90, status="cancelled", updated=at(hours=-3))
        c.put("Conference in Lisbon", at(days=21), all_day=True)

        d = self.drive
        d.put("Quarterly plan 2026", "application/vnd.google-apps.document",
              "iris@fabrikam.example", at(hours=-4),
              "Quarterly plan\n\nSection 2: the numbers for each team, and what changed.")
        d.put("Garden budget", "application/vnd.google-apps.document", ME,
              at(days=-3), "Seeds 40, tools 25, water 60.")

        t = self.tasks
        t.put_list("My Tasks")
        # Steve's own list (`modules/google_cloud/tasks.LIST_TITLE`).
        steve = t.put_list("Steve")
        t.put(steve, "Return the library books", due="2026-10-16")
        t.put(steve, "Pay the water bill", done=True)

    def close(self) -> None:
        self.mail.close()
        self.feeds.close()
        assert not getattr(self, "refused", []), f"Outlook reached unlisted: {self.refused}"


def wire(monkeypatch, world: World) -> None:
    """Connect the world at the services' own seams. The one part of this
    file that follows the code as it moves."""
    from agent import accounts, connectors
    from core import clock, db, modules
    from core.engines import ics, imap
    from modules.google_cloud import common as google_common
    from tests.fakes import google as fake_google
    from tests.fakes import outlook as fake_outlook

    monkeypatch.setattr(clock, "local_zone", lambda: ZONE)

    google_common._service.cache_clear()
    monkeypatch.setattr(google_common, "build", world.google.build)

    # Outlook, for every account's context: what it reaches beyond its
    # manifest's list is kept, for `close` to find empty.
    world.refused = fake_outlook.serve(monkeypatch, world.outlook)

    # The mailbox is added as the app adds an account. It takes the name and
    # the key of the stand-in it replaced, so the record's keys hold and only
    # real differences show.
    monkeypatch.setattr(connectors, "NETWORK", world.mail.network)
    # One network for the slot: the mail rig's sockets, and the feed's HTTP.
    monkeypatch.setattr(world.mail.network, "transport_for", world.feeds.network.transport_for)
    monkeypatch.setattr(imap, "SENT_WAIT_S", 0.2)
    ics.forget()
    conn = db.open_store()
    try:
        modules.register(conn)
        # Google, added and signed in as the app does it, under the key its
        # stand-in had until step 5's item 9, so the record's ids hold. One
        # name for its four classes: the stand-in's were "gmail" for mail
        # and "google" for the rest.
        fake_google.signed_in(conn, address=ME, name="gmail")
        # Classic Outlook, added, under the key and the name its stand-in had
        # until step 5's item 10.
        fake_outlook.add(conn)
        with pytest.MonkeyPatch.context() as keyed:
            keyed.setattr(modules, "new_key", lambda: "8c4d1477ffce")
            mailrig.add_account(conn, world.mail, address=persona.OWN_DOMAIN_EMAIL,
                                name="home")
        # The course site's link, added as the app adds one, with its "due
        # dates" box ticked: what the course feed's own key was.
        with pytest.MonkeyPatch.context() as keyed:
            keyed.setattr(modules, "new_key", lambda: "5e0b3a92c6d1")
            accounts.add(conn, module_id="ics", values={"link": COURSE_LINK, "due_dates": True},
                         name="Course site", resolve=lambda host, port: [feeds.PUBLIC_IP])
    finally:
        conn.close()
    # A Message-ID Steve writes holds the process's id and chance; the record
    # needs the same one each run.
    numbers = itertools.count(1)
    monkeypatch.setattr(email.utils, "make_msgid", lambda idstring=None, domain=None:
                        f"<golden-{next(numbers)}@{domain or 'golden.example'}>")
